"""Leitura em lotes e identificação explícita das fontes."""

import hashlib
import re
import unicodedata
from collections.abc import Iterator
from datetime import date, datetime
from itertools import islice
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def sha256(path: Path) -> str:
    with path.open("rb") as file:
        return hashlib.file_digest(file, "sha256").hexdigest()


def normalize_header(value: object) -> str:
    value = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Z0-9]+", "_", value.upper()).strip("_")


def family(path: Path) -> str:
    name = path.name.lower()
    if "trajetoria" in name:
        return "trajetoria"
    if "cursos" in name and path.suffix.lower() == ".csv":
        return "censo_cursos"
    if "ies" in name and path.suffix.lower() == ".csv":
        return "censo_ies"
    if "cpc" in name:
        return "cpc"
    if "enade" in name:
        return "enade"
    if "idade" in name:
        return "populacao"
    return "documentacao"


def workbook_metadata(workbook, kind: str) -> dict:
    """Preserva contexto e atualizações sem carregar as tabelas grandes."""
    metadata = {"sheet": workbook.worksheets[0].title, "workbook_sheets": workbook.sheetnames}
    supplementary = []
    for other in workbook.worksheets[1:]:
        preview = [
            [cell.isoformat() if isinstance(cell, date | datetime) else cell for cell in row]
            for row in other.iter_rows(max_row=min(other.max_row, 100), values_only=True)
        ]
        supplementary.append({
            "sheet": other.title, "rows": other.max_row,
            "preview": preview, "truncated": other.max_row > 100,
        })
        if kind == "populacao" and other.title.startswith("População -"):
            years = {str(cell) for row in preview[:10] for cell in row
                     if re.fullmatch(r"\d{4}", str(cell))}
            if len(years) == 1:
                metadata["reference_year"] = int(years.pop())
    metadata["supplementary_sheets"] = supplementary
    return metadata


def read_batches(path: Path, kind: str, batch_size: int) -> Iterator[tuple[pd.DataFrame, int]]:
    if path.suffix.lower() == ".csv":
        offset = 2
        for frame in pd.read_csv(
            path,
            sep=";",
            encoding="cp1252",
            dtype="string",
            keep_default_na=False,
            chunksize=batch_size,
            on_bad_lines="error",
        ):
            yield frame, offset
            offset += len(frame)
        return
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        metadata = workbook_metadata(workbook, kind)
        rows = sheet.iter_rows(values_only=True)
        header = None
        start = 0
        for index, row in enumerate(rows, 1):
            normalized = {normalize_header(cell) for cell in row if cell is not None}
            if kind != "trajetoria" or {"CO_CURSO", "NU_ANO_INGRESSO"} <= normalized:
                header = list(row)
                start = index + 1
                break
        if header is None:
            raise ValueError(f"Cabeçalho de dados não encontrado em {path.name}.")
        if any(cell is None for cell in header):
            raise ValueError(f"Cabeçalho incompleto em {path.name}.")
        data_count = 0
        while values := list(islice(rows, batch_size)):
            frame = pd.DataFrame(values, columns=header).astype("string")
            notes = frame.iloc[:, 0].str.startswith(("*", "Fonte:", "¹", "²", "³"), na=False)
            # CPC 2022 termina com uma contagem isolada na última coluna.
            # Não é um curso: todas as demais células estão vazias.
            last_number = pd.to_numeric(frame.iloc[:, -1], errors="coerce")
            candidate = (
                (kind == "cpc")
                & frame.iloc[:, :-1].isna().all(axis=1)
                & last_number.notna()
            ).fillna(False)
            blank = frame.isna().all(axis=1)
            preceding_data_count = (~notes & ~blank & ~candidate).cumsum() + data_count
            footer = candidate & last_number.eq(preceding_data_count).fillna(False)
            excluded = frame.loc[notes].iloc[:, 0].tolist()
            excluded += [f"Contagem isolada no rodapé: {v}" for v in frame.loc[footer].iloc[:, -1]]
            frame = frame.loc[~notes & ~blank & ~footer]
            data_count += len(frame)
            frame.attrs["excluded_notes"] = excluded
            frame.attrs["excluded_blank_rows"] = int(blank.sum())
            frame.attrs["source_rows"] = [start + int(i) for i in frame.index]
            frame.attrs["source_metadata"] = metadata
            yield frame, start
            start += len(values)
    finally:
        workbook.close()
